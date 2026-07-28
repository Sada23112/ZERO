"""Project ZERO — Document Reader Subsystem (Phase 4 Capability #4).

Reads PDF, Word, Markdown, Excel, CSV, JSON, YAML, & TOML documents and extracts structured elements.
"""

import json
import csv
from pathlib import Path
from typing import Dict, Any, Optional, List
from tools.base import BaseTool
from models.tool import ToolDefinition, ToolResult, ToolParameter
from zero_logging import logger

try:
    import yaml
    YAML_AVAILABLE = True
except ImportError:
    YAML_AVAILABLE = False

try:
    import toml
    TOML_AVAILABLE = True
except ImportError:
    TOML_AVAILABLE = False

try:
    from pypdf import PdfReader
    PYPDF_AVAILABLE = True
except ImportError:
    PYPDF_AVAILABLE = False

try:
    import docx
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    import openpyxl
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


class DocumentReaderTool(BaseTool):
    """Tool to read and parse PDF, Word, Markdown, Excel, CSV, JSON, YAML, and TOML documents."""

    def __init__(self, workspace_root: Optional[Path] = None):
        self.workspace_root = (workspace_root or Path.cwd()).resolve()

    @property
    def name(self) -> str:
        return "document_reader"

    @property
    def description(self) -> str:
        return "Read and extract structured content from PDF, Word (.docx), Excel (.xlsx), CSV, JSON, YAML, TOML, and Markdown files."

    def get_definition(self) -> ToolDefinition:
        return ToolDefinition(
            name=self.name,
            description=self.description,
            parameters={
                "path": ToolParameter(type="string", description="Document file path to read")
            }
        )

    async def execute(self, call_id: str, arguments: Dict[str, Any]) -> ToolResult:
        file_str = arguments.get("path")
        if not file_str:
            return ToolResult(call_id=call_id, tool_name=self.name, success=False, error="Argument 'path' is required.")

        target_p = Path(file_str).resolve()
        if not target_p.exists():
            return ToolResult(call_id=call_id, tool_name=self.name, success=False, error=f"File not found: {file_str}")

        ext = target_p.suffix.lower()

        try:
            # 1. JSON / YAML / TOML
            if ext == ".json":
                data = json.loads(target_p.read_text(encoding="utf-8"))
                return ToolResult(call_id=call_id, tool_name=self.name, success=True, output=f"[JSON Document]\n{json.dumps(data, indent=2)[:3000]}")

            elif ext in [".yaml", ".yml"]:
                if YAML_AVAILABLE:
                    data = yaml.safe_load(target_p.read_text(encoding="utf-8"))
                    return ToolResult(call_id=call_id, tool_name=self.name, success=True, output=f"[YAML Document]\n{json.dumps(data, indent=2)[:3000]}")
                return ToolResult(call_id=call_id, tool_name=self.name, success=True, output=target_p.read_text(encoding="utf-8")[:3000])

            elif ext == ".toml":
                if TOML_AVAILABLE:
                    data = toml.loads(target_p.read_text(encoding="utf-8"))
                    return ToolResult(call_id=call_id, tool_name=self.name, success=True, output=f"[TOML Document]\n{json.dumps(data, indent=2)[:3000]}")
                return ToolResult(call_id=call_id, tool_name=self.name, success=True, output=target_p.read_text(encoding="utf-8")[:3000])

            # 2. CSV / Excel
            elif ext == ".csv":
                rows = []
                with open(target_p, "r", encoding="utf-8", errors="ignore") as f:
                    reader = csv.reader(f)
                    for r in list(reader)[:20]:
                        rows.append(" | ".join(r))
                return ToolResult(call_id=call_id, tool_name=self.name, success=True, output=f"[CSV Document ({len(rows)} preview rows)]:\n" + "\n".join(rows))

            elif ext in [".xlsx", ".xls"]:
                if OPENPYXL_AVAILABLE:
                    wb = openpyxl.load_workbook(target_p, read_only=True)
                    sheet = wb.active
                    rows = []
                    for row in list(sheet.iter_rows(values_only=True))[:20]:
                        rows.append(" | ".join([str(val) for val in row if val is not None]))
                    return ToolResult(call_id=call_id, tool_name=self.name, success=True, output=f"[Excel Sheet '{sheet.title}']:\n" + "\n".join(rows))

            # 3. PDF Document
            elif ext == ".pdf":
                if PYPDF_AVAILABLE:
                    reader = PdfReader(str(target_p))
                    pages_text = []
                    for i, page in enumerate(reader.pages[:5]):
                        pages_text.append(f"--- Page {i+1} ---\n" + (page.extract_text() or ""))
                    return ToolResult(call_id=call_id, tool_name=self.name, success=True, output=f"[PDF Document ({len(reader.pages)} pages)]:\n" + "\n".join(pages_text))

            # 4. Word Document (.docx)
            elif ext == ".docx":
                if DOCX_AVAILABLE:
                    doc = docx.Document(str(target_p))
                    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
                    return ToolResult(call_id=call_id, tool_name=self.name, success=True, output=f"[Word Document]:\n" + "\n".join(paragraphs[:30]))

            # Fallback Text / Markdown
            text_content = target_p.read_text(encoding="utf-8", errors="ignore")
            return ToolResult(call_id=call_id, tool_name=self.name, success=True, output=f"[{ext.upper()} Document]:\n{text_content[:3000]}")

        except Exception as err:
            logger.error(f"Error reading document {file_str}: {err}")
            return ToolResult(call_id=call_id, tool_name=self.name, success=False, error=str(err))
